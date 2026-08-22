from datetime import datetime, timedelta
from io import BytesIO

from django.conf import settings
from django.http import HttpResponse
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment, Border, Side
from rest_framework import status
from rest_framework.decorators import api_view, permission_classes
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response

from ..models import Attendance, AttendanceSession, Batch, Student
from ..permissions import IsAdminRole


SUBJECT_ABBREVIATIONS = {
    "human values and ethics": "HVT",
}
EXCEL_UNAVAILABLE = "-"


def _staff_short_name(staff_name):
    words = [word for word in staff_name.split() if word]
    if len(words) >= 2:
        return f"{words[-1][0]}{words[0][0]}".upper()
    return words[0][:2].upper() if words else "N/C"


def _subject_short_name(subject_name):
    abbreviation = SUBJECT_ABBREVIATIONS.get(subject_name.strip().lower())
    if abbreviation:
        return abbreviation

    ignored_words = {"a", "an", "and", "of", "the"}
    initials = [
        word[0]
        for word in subject_name.split()
        if word and word.lower() not in ignored_words
    ]
    return "".join(initials).upper() or "N/C"


def _weekly_report_response(request):
    class_name = request.query_params.get("class_name", "").strip()
    batch_id = request.query_params.get("batch", "").strip()
    start_date = request.query_params.get("start_date", "").strip()
    end_date = request.query_params.get("end_date", "").strip()

    if not class_name or not batch_id or not start_date or not end_date:
        return Response({
            "message": "class_name, batch, start_date and end_date are required"
        }, status=status.HTTP_400_BAD_REQUEST)

    try:
        batch = Batch.objects.get(id=batch_id, is_active=True)
    except (Batch.DoesNotExist, ValueError):
        return Response({"message": "Active batch not found"}, status=status.HTTP_404_NOT_FOUND)

    try:
        start = datetime.strptime(start_date, "%Y-%m-%d").date()
        end = datetime.strptime(end_date, "%Y-%m-%d").date()
    except ValueError:
        return Response({"message": "Dates must use YYYY-MM-DD format"}, status=status.HTTP_400_BAD_REQUEST)
    if start > end:
        return Response({"message": "start_date must be before end_date"}, status=status.HTTP_400_BAD_REQUEST)

    students = list(
        Student.objects.filter(class_name=class_name, batch=batch)
        .order_by("register_no")
    )
    sessions = list(
        AttendanceSession.objects.filter(
            class_name=class_name,
            date__range=(start, end),
            status="COMPLETED",
        )
        .select_related("staff", "subject")
        .order_by("date", "period", "id")
    )

    sessions_by_slot = {}
    for session in sessions:
        sessions_by_slot.setdefault((session.date, session.period), session)

    attendance_by_slot_student = {
        (record.session.date.isoformat(), record.session.period, record.student_id): record.status
        for record in Attendance.objects.filter(
            session__in=sessions,
            student__in=students,
        ).select_related("session")
    }

    slots = []
    current_date = start
    while current_date <= end:
        for period in range(1, 9):
            session = sessions_by_slot.get((current_date, period))
            slots.append({
                "date": current_date.isoformat(),
                "day": current_date.strftime("%A"),
                "period": period,
                "staff": session.staff.staff_name if session else "N/C",
                "subject_code": session.subject.subject_code if session else "N/C",
                "subject": session.subject.subject_name if session else "N/C",
                "session_id": session.id if session else None,
            })
        current_date += timedelta(days=1)

    student_rows = []
    counted_slots = len(sessions_by_slot)
    for student in students:
        attendance = {}
        present = 0
        absent = 0
        for slot in slots:
            key = (slot["date"], slot["period"], student.id)
            value = attendance_by_slot_student.get(key, "M" if slot["session_id"] else "N/C")
            attendance[f"{slot['date']}_P{slot['period']}"] = value
            if value == "Present":
                present += 1
            elif value == "Absent":
                absent += 1

        student_rows.append({
            "register_no": student.register_no,
            "student_name": student.student_name,
            "attendance": attendance,
            "total_classes": counted_slots,
            "present": present,
            "absent": absent,
            "missing": counted_slots - present - absent,
            "attendance_percentage": round((present / (present + absent)) * 100, 2) if present + absent else 0,
        })

    return Response({
        "class_name": class_name,
        "batch": {"id": batch.id, "name": batch.name},
        "start_date": start.isoformat(),
        "end_date": end.isoformat(),
        "slots": slots,
        "students": student_rows,
    })


@api_view(["GET"])
@permission_classes([IsAdminRole])
def weekly_report(request):
    return _weekly_report_response(request)


@api_view(["GET"])
@permission_classes([IsAdminRole])
def weekly_report_excel(request):
    report = _weekly_report_response(request)
    if report.status_code != status.HTTP_200_OK:
        return report

    report_data = report.data
    slots = report_data["slots"]
    dates = []
    for slot in slots:
        if slot["date"] not in dates:
            dates.append(slot["date"])

    if len(dates) > 6:
        return Response({
            "message": "The Excel template supports a maximum of 6 dates"
        }, status=status.HTTP_400_BAD_REQUEST)

    template_path = settings.BASE_DIR / "III IT B WEEKLY ATT.xlsx"
    workbook = load_workbook(template_path)
    worksheet = workbook["Master"]

    for sheet in list(workbook.worksheets):
        if sheet.title != "Master":
            workbook.remove(sheet)
    workbook.active = workbook.index(worksheet)

    for merged_range in list(worksheet.merged_cells.ranges):
        if merged_range.min_row in (8, 9) and merged_range.max_row <= 9:
            if merged_range.min_col >= 5 and merged_range.max_col <= 50:
                worksheet.unmerge_cells(str(merged_range))

    header_edge = Side(style="medium", color="000000")
    no_header_edge = Side(style=None)
    for row in (8, 9):
        for column in range(5, 51):
            is_date_start = (column - 5) % 8 == 0
            is_date_end = (column - 5) % 8 == 7
            worksheet.cell(row, column).border = Border(
                left=header_edge if is_date_start else no_header_edge,
                right=header_edge if is_date_end else no_header_edge,
                top=header_edge,
                bottom=header_edge,
            )
            worksheet.cell(row, column).alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
                shrink_to_fit=True,
            )

    worksheet["A3"] = "Weekly Attendance Report"
    worksheet["A4"] = (
        f"Attendance From: {report_data['start_date']} "
        f"to {report_data['end_date']}"
    )
    worksheet["U4"] = f"BATCH: {report_data['batch']['name']}"
    worksheet["AL4"] = f"CLASS: {report_data['class_name']}"

    for column in range(5, 51):
        worksheet.cell(8, column).value = EXCEL_UNAVAILABLE
        worksheet.cell(9, column).value = EXCEL_UNAVAILABLE

    slot_by_key = {
        (slot["date"], slot["period"]): slot
        for slot in slots
    }
    for date_index, report_date in enumerate(dates):
        first_column = 5 + date_index * 8
        worksheet.cell(5, first_column).value = report_date
        worksheet.cell(6, first_column).value = datetime.strptime(
            report_date, "%Y-%m-%d"
        ).strftime("%A").upper()
        for period in range(1, 9):
            column = first_column + period - 1
            slot = slot_by_key[(report_date, period)]
            worksheet.cell(7, column).value = period
            worksheet.cell(8, column).value = (
                _subject_short_name(slot["subject"])
                if slot["session_id"] else EXCEL_UNAVAILABLE
            )
            worksheet.cell(9, column).value = (
                _staff_short_name(slot["staff"])
                if slot["session_id"] else EXCEL_UNAVAILABLE
            )

    if len(report_data["students"]) > worksheet.max_row - 9:
        return Response({
            "message": "The Excel template supports a maximum of 71 students"
        }, status=status.HTTP_400_BAD_REQUEST)

    for row in range(10, worksheet.max_row + 1):
        for column in range(1, 56):
            if column != 3 and not isinstance(worksheet.cell(row, column), MergedCell):
                worksheet.cell(row, column).value = None
        if not isinstance(worksheet.cell(row, 3), MergedCell):
            worksheet.cell(row, 3).value = None

    for row_number, student in enumerate(report_data["students"], start=10):
        worksheet.cell(row_number, 1).value = student["register_no"]
        worksheet.cell(row_number, 2).value = student["student_name"]
        worksheet.cell(row_number, 3).value = None
        worksheet.cell(row_number, 4).value = 0

        present = 0
        for date_index, report_date in enumerate(dates):
            first_column = 5 + date_index * 8
            for period in range(1, 9):
                key = f"{report_date}_P{period}"
                value = student["attendance"].get(key, "N/C")
                if value == "Present":
                    cell_value = 1
                    present += 1
                elif value == "Absent":
                    cell_value = "a"
                elif value == "N/C":
                    cell_value = EXCEL_UNAVAILABLE
                else:
                    cell_value = value
                worksheet.cell(row_number, first_column + period - 1).value = cell_value

        worksheet.cell(row_number, 53).value = present
        worksheet.cell(row_number, 54).value = present
        worksheet.cell(row_number, 55).value = (
            round((present / (present + student["absent"])) * 100, 2)
            if present + student["absent"] else 0
        )

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = (
        f"attachment; filename=weekly_report_{report_data['batch']['name']}.xlsx"
    )
    return response
