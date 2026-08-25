from io import BytesIO

from openpyxl import Workbook


def generate_staff_import_result_excel(
    created_staff,
    ignored_staff,
    failed_staff,
):

    workbook = Workbook()

    # ========================================================
    # CREATED
    # ========================================================

    worksheet = workbook.active

    worksheet.title = "Created Staff"

    worksheet.append([
        "Staff ID",
        "Staff Name",
        "Username",
        "Temporary Password",
    ])

    for staff in created_staff:

        worksheet.append([
            staff.get(
                "staff_id",
                "",
            ),

            staff.get(
                "staff_name",
                "",
            ),

            staff.get(
                "username",
                "",
            ),

            staff.get(
                "temporary_password",
                "",
            ),
        ])

    # ========================================================
    # IGNORED
    # ========================================================

    worksheet = workbook.create_sheet(
        "Ignored Staff"
    )

    worksheet.append([
        "Staff ID",
        "Staff Name",
        "Reason",
    ])

    for staff in ignored_staff:

        worksheet.append([
            staff.get(
                "staff_id",
                "",
            ),

            staff.get(
                "staff_name",
                "",
            ),

            staff.get(
                "reason",
                "",
            ),
        ])

    # ========================================================
    # FAILED
    # ========================================================

    worksheet = workbook.create_sheet(
        "Failed Staff"
    )
    worksheet.append([
        "Sheet",
        "Row",
        "Staff ID",
        "Reason",
    ])
    for staff in failed_staff:
        worksheet.append([
            staff.get(
                "sheet",
                "",
            ),
            staff.get(
                "row",
                "",
            ),
            staff.get(
                "staff_id",
                "",
            ),
            staff.get(
                "reason",
                "",
            ),
        ])
    # ========================================================
    # COLUMN WIDTH
    # ========================================================
    for worksheet in workbook.worksheets:
        for column in worksheet.columns:
            maximum = 0
            letter = (
                column[0]
                .column_letter
            )
            for cell in column:
                if cell.value is not None:
                    maximum = max(
                        maximum,
                        len(
                            str(
                                cell.value
                            )
                        ),
                    )
            worksheet.column_dimensions[
                letter
            ].width = min(
                maximum + 2,
                50,
            )
    # ========================================================
    # RETURN
    # ========================================================
    output = BytesIO()

    workbook.save(output)

    output.seek(0)

    return output