from io import BytesIO

from django.contrib.auth.hashers import make_password
from django.db import transaction

from openpyxl import Workbook

from ..models import (
    User,
    Student,
    Batch,
)

from ..utils import (
    generate_username,
    generate_password,
)

from .file_parser_service import (
    parse_file,
)

from .normalization_service import (
    normalize_batch,
)


# ============================================================
# VALIDATION
# ============================================================

def validate_student_record(record):

    errors = []

    if not record["register_no"]:
        errors.append("Register No is missing")

    if not record["student_name"]:
        errors.append("Student Name is missing")

    if not record["class_name"]:
        errors.append("Class is missing")

    if not record["batch"]:
        errors.append("Batch is missing")

    return errors


# ============================================================
# BATCH CACHE
# ============================================================

def build_batch_cache():

    batches = Batch.objects.filter(
        is_active=True
    )

    return {
        normalize_batch(batch.name): batch
        for batch in batches
    }


# ============================================================
# IMPORT
# ============================================================

@transaction.atomic
def import_students_from_file(
    uploaded_file,
    admin_user=None,
):

    records = parse_file(
        uploaded_file
    )

    created_students = []
    ignored_students = []
    failed_students = []

    if not records:

        return {
            "created_students": [],
            "ignored_students": [],
            "failed_students": [],
        }

    # --------------------------------------------------------
    # Batch cache
    # --------------------------------------------------------

    batch_cache = build_batch_cache()

    # --------------------------------------------------------
    # Existing register numbers
    #
    # One DB query instead of one query per row.
    # --------------------------------------------------------

    register_numbers = {
        record["register_no"]
        for record in records
        if record["register_no"]
    }

    existing_students = {
        student.register_no: student
        for student in Student.objects.filter(
            register_no__in=register_numbers
        ).select_related("user")
    }

    # --------------------------------------------------------
    # Detect duplicate register numbers inside upload
    # --------------------------------------------------------

    seen_register_numbers = set()

    students_to_create = []

    # --------------------------------------------------------
    # First pass
    # --------------------------------------------------------

    for record in records:

        errors = validate_student_record(
            record
        )

        if errors:

            failed_students.append({
                "sheet": record.get("_sheet"),
                "row": record.get("_row"),
                "register_no": record.get(
                    "register_no",
                    "",
                ),
                "reason": "; ".join(errors),
            })

            continue

        register_no = record["register_no"]

        # Existing DB student.
        if register_no in existing_students:

            existing = existing_students[
                register_no
            ]

            ignored_students.append({
                "register_no": existing.register_no,
                "student_name": existing.student_name,
                "reason": "Register No already exists",
            })

            continue

        # Duplicate inside uploaded file.
        if register_no in seen_register_numbers:

            ignored_students.append({
                "register_no": register_no,
                "student_name": record["student_name"],
                "reason": (
                    "Duplicate Register No "
                    "inside uploaded file"
                ),
            })

            continue

        seen_register_numbers.add(
            register_no
        )

        # Batch lookup from cache.
        batch = batch_cache.get(
            normalize_batch(
                record["batch"]
            )
        )

        if not batch:

            failed_students.append({
                "sheet": record.get("_sheet"),
                "row": record.get("_row"),
                "register_no": register_no,
                "reason": (
                    f"Active batch "
                    f"'{record['batch']}' "
                    f"not found"
                ),
            })

            continue

        students_to_create.append({
            "record": record,
            "batch": batch,
        })

    # --------------------------------------------------------
    # Create users + students
    # --------------------------------------------------------

    for item in students_to_create:

        record = item["record"]
        batch = item["batch"]

        raw_password = generate_password()
        username = generate_username("STU")

        user = User.objects.create(
            username=username,
            password=make_password(
                raw_password
            ),
            role="STUDENT",
        )

        student = Student.objects.create(
            user=user,
            student_name=record["student_name"],
            register_no=record["register_no"],
            class_name=record["class_name"],
            batch=batch,
        )

        created_students.append({
            "register_no": student.register_no,
            "student_name": student.student_name,
            "class_name": student.class_name,
            "batch": batch.name,
            "username": username,
            "temporary_password": raw_password,
        })

    return {
        "created_students": created_students,
        "ignored_students": ignored_students,
        "failed_students": failed_students,
    }


# ============================================================
# BACKWARD COMPATIBILITY
# ============================================================

def import_students_from_excel(
    uploaded_file,
    admin_user=None,
):

    return import_students_from_file(
        uploaded_file,
        admin_user=admin_user,
    )


# ============================================================
# RESULT EXCEL
# ============================================================

def generate_credentials_excel(
    created_students,
    ignored_students=None,
    failed_students=None,
):

    ignored_students = ignored_students or []
    failed_students = failed_students or []

    workbook = Workbook()

    # ========================================================
    # CREATED
    # ========================================================

    worksheet = workbook.active
    worksheet.title = "Created Students"

    worksheet.append([
        "Reg No",
        "Student Name",
        "Class",
        "Batch",
        "Username",
        "Temporary Password",
    ])

    for student in created_students:

        worksheet.append([
            student["register_no"],
            student["student_name"],
            student["class_name"],
            student["batch"],
            student["username"],
            student["temporary_password"],
        ])

    # ========================================================
    # IGNORED
    # ========================================================

    ignored_sheet = workbook.create_sheet(
        "Ignored Students"
    )

    ignored_sheet.append([
        "Reg No",
        "Student Name",
        "Reason",
    ])

    for student in ignored_students:

        ignored_sheet.append([
            student.get("register_no", ""),
            student.get("student_name", ""),
            student.get("reason", ""),
        ])

    # ========================================================
    # FAILED
    # ========================================================

    failed_sheet = workbook.create_sheet(
        "Failed Students"
    )

    failed_sheet.append([
        "Sheet",
        "Row",
        "Reg No",
        "Reason",
    ])

    for student in failed_students:

        failed_sheet.append([
            student.get("sheet", ""),
            student.get("row", ""),
            student.get("register_no", ""),
            student.get("reason", ""),
        ])

    # ========================================================
    # COLUMN WIDTHS
    # ========================================================

    widths = {
        "A": 20,
        "B": 30,
        "C": 20,
        "D": 18,
        "E": 22,
        "F": 25,
    }

    for sheet in workbook.worksheets:

        for column, width in widths.items():
            sheet.column_dimensions[
                column
            ].width = width

    output = BytesIO()

    workbook.save(output)

    output.seek(0)

    return output