from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from .normalization_service import (
    REQUIRED_FIELDS,
    STAFF_ID_ALIASES,
    STAFF_NAME_ALIASES,
    map_columns,
    normalize_student_record,
)


SUPPORTED_EXTENSIONS = {
    ".xlsx",
    ".csv",
    ".pdf",
}


# ============================================================
# COMMON
# ============================================================

def _has_data(record):
    return any(
        value not in (None, "")
        for value in record.values()
    )


def _map_row(raw_record, sheet, row):

    record = dict(raw_record)

    record["_sheet"] = sheet
    record["_row"] = row

    return normalize_student_record(record)


# ============================================================
# EXCEL
# ============================================================

def parse_excel(uploaded_file):

    workbook = load_workbook(
        uploaded_file,
        read_only=True,
        data_only=True,
    )

    records = []

    for worksheet in workbook.worksheets:

        if worksheet.max_row == 0:
            continue

        header_row = None
        header_mapping = None

        # Find header in first 20 rows.
        for row_number in range(
            1,
            min(worksheet.max_row, 20) + 1
        ):

            headers = [
                cell.value
                for cell in worksheet[row_number]
            ]

            try:
                mapping = map_columns(headers)
            except ValueError:
                continue

            header_row = row_number
            header_mapping = mapping
            break

        if header_row is None:
            continue

        # Canonical field -> Excel column number
        columns = {}

        for column_number, cell in enumerate(
            worksheet[header_row],
            start=1,
        ):

            field = header_mapping.get(
                cell.value
            )

            if field:
                columns[field] = column_number

        for row_number in range(
            header_row + 1,
            worksheet.max_row + 1,
        ):

            raw_record = {}

            for field, column_number in columns.items():

                raw_record[field] = worksheet.cell(
                    row_number,
                    column_number,
                ).value

            if not _has_data(raw_record):
                continue

            records.append(
                _map_row(
                    raw_record,
                    worksheet.title,
                    row_number,
                )
            )

    return records


# ============================================================
# CSV
# ============================================================

def parse_csv(uploaded_file):

    uploaded_file.seek(0)

    try:

        df = pd.read_csv(
            uploaded_file,
            encoding="utf-8-sig",
        )

    except UnicodeDecodeError:

        uploaded_file.seek(0)

        df = pd.read_csv(
            uploaded_file,
            encoding="latin-1",
        )

    if df.empty:
        return []

    mapping = map_columns(
        df.columns
    )

    df = df.rename(
        columns=mapping
    )

    records = []

    for index, row in df.iterrows():

        raw_record = {
            field: row.get(field)
            for field in REQUIRED_FIELDS
        }

        if not any(
            pd.notna(value)
            and str(value).strip()
            for value in raw_record.values()
        ):
            continue

        records.append(
            _map_row(
                raw_record,
                "CSV",
                index + 2,
            )
        )

    return records


# ============================================================
# PDF
# ============================================================

def parse_pdf(uploaded_file):

    try:
        import pdfplumber
    except ImportError:
        raise ValueError(
            "PDF support requires pdfplumber. "
            "Install it with: pip install pdfplumber"
        )

    records = []

    with pdfplumber.open(uploaded_file) as pdf:

        for page_number, page in enumerate(
            pdf.pages,
            start=1,
        ):

            tables = page.extract_tables()

            for table in tables:

                if not table:
                    continue

                header = table[0]

                if not header:
                    continue

                try:
                    mapping = map_columns(header)
                except ValueError:
                    continue

                indexes = {}

                for index, column in enumerate(header):

                    field = mapping.get(column)

                    if field:
                        indexes[field] = index

                for row_number, row in enumerate(
                    table[1:],
                    start=2,
                ):

                    if not row:
                        continue

                    raw_record = {}

                    for field, index in indexes.items():

                        raw_record[field] = (
                            row[index]
                            if index < len(row)
                            else ""
                        )

                    if not _has_data(raw_record):
                        continue

                    records.append(
                        _map_row(
                            raw_record,
                            f"PDF Page {page_number}",
                            row_number,
                        )
                    )

    return records


# ============================================================
# MAIN
# ============================================================

def parse_file(uploaded_file):

    filename = Path(
        uploaded_file.name
    ).suffix.lower()

    if filename == ".xlsx":
        return parse_excel(uploaded_file)

    if filename == ".csv":
        return parse_csv(uploaded_file)

    if filename == ".pdf":
        return parse_pdf(uploaded_file)

    raise ValueError(
        "Unsupported file format. "
        "Supported formats: .xlsx, .csv, .pdf"
    )
import csv
import io

from openpyxl import load_workbook

from .normalization_service import (
    detect_staff_field,
)


# ============================================================
# STAFF HEADER MAP
# ============================================================

def build_staff_header_map(headers):

    header_map = {}

    for index, header in enumerate(headers):

        field = detect_staff_field(header)

        if field and field not in header_map:
            header_map[field] = index

    required_fields = {
        "staff_name",
        "staff_id",
    }

    missing = (
        required_fields
        - set(header_map.keys())
    )

    if missing:

        raise ValueError(
            "Required Staff columns missing: "
            + ", ".join(sorted(missing))
        )

    return header_map


# ============================================================
# XLSX
# ============================================================

def parse_staff_xlsx(uploaded_file):

    workbook = load_workbook(
        uploaded_file,
        read_only=True,
        data_only=True,
    )

    rows = []

    for worksheet in workbook.worksheets:

        if worksheet.max_row == 0:
            continue

        # ----------------------------------------------------
        # FIND HEADER
        # ----------------------------------------------------

        header_row = None
        header_map = None

        for row_number, row in enumerate(
            worksheet.iter_rows(
                min_row=1,
                max_row=min(
                    worksheet.max_row,
                    20,
                ),
                values_only=True,
            ),
            start=1,
        ):

            try:

                candidate_map = (
                    build_staff_header_map(row)
                )

                header_row = row_number
                header_map = candidate_map

                break

            except ValueError:
                continue

        if header_map is None:
            continue

        # ----------------------------------------------------
        # READ DATA
        # ----------------------------------------------------

        for row_number, row in enumerate(
            worksheet.iter_rows(
                min_row=header_row + 1,
                values_only=True,
            ),
            start=header_row + 1,
        ):

            staff_name = ""

            staff_id = ""

            name_index = header_map[
                "staff_name"
            ]

            id_index = header_map[
                "staff_id"
            ]

            if name_index < len(row):
                staff_name = row[
                    name_index
                ]

            if id_index < len(row):
                staff_id = row[
                    id_index
                ]

            # Empty row
            if not (
                staff_name
                or staff_id
            ):
                continue

            rows.append({
                "staff_name": staff_name,
                "staff_id": staff_id,
                "_sheet": worksheet.title,
                "_row": row_number,
            })

    workbook.close()

    return rows


# ============================================================
# CSV
# ============================================================

def parse_staff_csv(uploaded_file):

    content = uploaded_file.read()

    if isinstance(content, bytes):

        content = content.decode(
            "utf-8-sig",
            errors="replace",
        )

    reader = csv.reader(
        io.StringIO(content)
    )

    all_rows = list(reader)

    if not all_rows:
        return []

    header_row = None
    header_map = None

    # --------------------------------------------------------
    # FIND HEADER
    # --------------------------------------------------------

    for index, row in enumerate(
        all_rows[:20]
    ):

        try:

            candidate_map = (
                build_staff_header_map(row)
            )

            header_row = index
            header_map = candidate_map

            break

        except ValueError:
            continue

    if header_map is None:

        raise ValueError(
            "Could not identify Staff Name "
            "and Staff ID headers."
        )

    rows = []

    # --------------------------------------------------------
    # DATA
    # --------------------------------------------------------

    for row_number, row in enumerate(
        all_rows[
            header_row + 1:
        ],
        start=header_row + 2,
    ):

        staff_name = ""

        staff_id = ""

        name_index = header_map[
            "staff_name"
        ]

        id_index = header_map[
            "staff_id"
        ]

        if name_index < len(row):
            staff_name = row[
                name_index
            ]

        if id_index < len(row):
            staff_id = row[
                id_index
            ]

        if not (
            staff_name
            or staff_id
        ):
            continue

        rows.append({
            "staff_name": staff_name,
            "staff_id": staff_id,
            "_sheet": "CSV",
            "_row": row_number,
        })

    return rows


# ============================================================
# PDF
# ============================================================

def parse_staff_pdf(uploaded_file):

    try:

        import pdfplumber

    except ImportError:

        raise ValueError(
            "PDF support requires pdfplumber."
        )

    rows = []

    with pdfplumber.open(
        uploaded_file
    ) as pdf:

        for page_number, page in enumerate(
            pdf.pages,
            start=1,
        ):

            tables = page.extract_tables()

            for table in tables:

                if not table:
                    continue

                header_row = None
                header_map = None

                # --------------------------------------------
                # FIND HEADER
                # --------------------------------------------

                for index, row in enumerate(
                    table[:20]
                ):

                    if not row:
                        continue

                    try:

                        candidate_map = (
                            build_staff_header_map(
                                row
                            )
                        )

                        header_row = index
                        header_map = candidate_map

                        break

                    except ValueError:
                        continue

                if header_map is None:
                    continue

                # --------------------------------------------
                # DATA
                # --------------------------------------------

                for row_number, row in enumerate(
                    table[
                        header_row + 1:
                    ],
                    start=header_row + 2,
                ):

                    if not row:
                        continue

                    name_index = header_map[
                        "staff_name"
                    ]

                    id_index = header_map[
                        "staff_id"
                    ]

                    staff_name = ""

                    staff_id = ""

                    if name_index < len(row):
                        staff_name = row[
                            name_index
                        ]

                    if id_index < len(row):
                        staff_id = row[
                            id_index
                        ]

                    if not (
                        staff_name
                        or staff_id
                    ):
                        continue

                    rows.append({
                        "staff_name": staff_name,
                        "staff_id": staff_id,
                        "_sheet": (
                            f"PDF Page {page_number}"
                        ),
                        "_row": row_number,
                    })

    return rows



# ============================================================
# MAIN STAFF PARSER
# ============================================================

def parse_staff_file(uploaded_file):

    filename = (
        uploaded_file.name
        .lower()
        .strip()
    )

    if filename.endswith(".xlsx"):

        return parse_staff_xlsx(
            uploaded_file
        )

    if filename.endswith(".csv"):

        return parse_staff_csv(
            uploaded_file
        )

    if filename.endswith(".pdf"):

        return parse_staff_pdf(
            uploaded_file
        )

    raise ValueError(
        "Unsupported file format. "
        "Supported formats: XLSX, CSV, PDF."
    )